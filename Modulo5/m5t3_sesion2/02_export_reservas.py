"""
02_export_reservas.py  —  Del triángulo chain-ladder al IBNR, LLENANDO una plantilla
====================================================================================
Módulo 5 · Tema 2 · Parte 3 · VIERNES (reservas)

Hoy NO nacemos el archivo: la empresa / CNSF nos da una CÉDULA fija y la ALIMENTAMOS.
Esa es la diferencia con el miércoles:
    miércoles -> xlsxwriter CREA el Excel desde cero
    viernes   -> openpyxl   LLENA una plantilla que ya existe (load_workbook)

Calculamos en Python (chain-ladder + Bornhuetter-Ferguson) y escribimos en las celdas
fijas de la plantilla. Dejamos el IBNR como FÓRMULAS vivas y una hoja de CONTROL
(reconciliación Python vs Excel).

Se corre:  python 02_export_reservas.py   ->   reservas_cedula.xlsx
(requiere  plantilla_reservas.xlsx  en la misma carpeta)
"""
import numpy as np
import pandas as pd
import chainladder as cl
import openpyxl
from openpyxl.utils import get_column_letter as L

# ──────────────────────────────────────────────────────────────────────────
# 1 · Triángulo + chain-ladder (esto ya lo conocen; aquí el foco es EXPORTAR)
# ──────────────────────────────────────────────────────────────────────────
tri   = cl.load_sample("raa")                               # triángulo de ejemplo
frame = tri.to_frame(origin_as_datetime=False)              # acumulado (10×10)
anios = [int(getattr(a, "year", a)) for a in frame.index]   # 1981..1990
AN    = len(anios)                                          # 10

ldf  = cl.Development().fit(tri).ldf_.to_frame(origin_as_datetime=False).values.flatten()
ldf  = [float(x) for x in ldf]                              # 9 factores de desarrollo
TAIL = 1.0

M = frame.values                                            # matriz del triángulo
ult_cl, ibnr_cl, cdf, last = [], [], [], []
for j in range(AN):                                         # j=0 (1981) ... j=9 (1990)
    lc   = M[j, AN - 1 - j]                                 # último acumulado (diagonal)
    c    = float(np.prod(ldf[AN - 1 - j:])) * TAIL          # CDF = Π factores restantes
    u    = lc * c
    last.append(lc); cdf.append(c); ult_cl.append(u); ibnr_cl.append(u - lc)

IBNR_TOTAL = float(np.nansum(ibnr_cl))
print("IBNR total (Python):", round(IBNR_TOTAL))
print("IBNR total (chainladder):", round(float(cl.Chainladder().fit(tri).ibnr_.sum())))

# ──────────────────────────────────────────────────────────────────────────
# 2 · Bornhuetter-Ferguson (a priori) — para comparar métodos
#     IBNR_BF = ultimate_a_priori × (1 − 1/CDF)        (lo no desarrollado del a priori)
# ──────────────────────────────────────────────────────────────────────────
PRIMA = np.linspace(18000, 27000, AN)          # prima devengada por año (supuesto)
ELR   = 0.75                                    # loss ratio esperado (supuesto)
apriori_ult = PRIMA * ELR
ibnr_bf = [apriori_ult[j] * (1 - 1 / cdf[j]) for j in range(AN)]

# ──────────────────────────────────────────────────────────────────────────
# 3 · LLENAR la plantilla con openpyxl (celdas fijas + fórmulas vivas)
# ──────────────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook("plantilla_reservas.xlsx")      # ← abrimos lo que ya existe
ws = wb["Cedula"]
MONEY, FAC = "#,##0", "0.000"

# 3a · triángulo (valores donde hay dato; NaN se deja vacío)
for i in range(AN):
    ws.cell(6 + i, 1, anios[i])
    for d in range(AN):
        v = M[i, d]
        if not np.isnan(v):
            c = ws.cell(6 + i, 2 + d, float(v)); c.number_format = MONEY

# 3b · factores de desarrollo + tail
for d in range(AN - 1):
    # OJO: guardamos el factor con PRECISIÓN COMPLETA (el formato solo lo MUESTRA a 3 decimales).
    # Si redondeáramos el valor, el CDF del Excel no cuadraría con Python (la reconciliación fallaría).
    c = ws.cell(19, 2 + d, float(ldf[d])); c.number_format = FAC
c = ws.cell(19, 1 + AN, TAIL); c.number_format = FAC         # tail en col K

# 3c · tabla de reservas: B=último acumulado, C=CDF, D=ultimate, E=IBNR  (FÓRMULAS vivas)
for j in range(AN):
    r = 23 + j
    ws.cell(r, 1, anios[j])
    diag = f"{L(11 - j)}{6 + j}"                              # celda diagonal del triángulo
    ws.cell(r, 2, f"={diag}").number_format = MONEY           # último acumulado (vive del triángulo)
    if j == 0:
        ws.cell(r, 3, "=K19").number_format = FAC             # CDF (año desarrollado = tail)
    else:
        rng = f"{L(11 - j)}19:J19"                            # factores restantes
        ws.cell(r, 3, f"=PRODUCT({rng})*K19").number_format = FAC
    ws.cell(r, 4, f"=B{r}*C{r}").number_format = MONEY        # ultimate = acumulado × CDF
    ws.cell(r, 5, f"=D{r}-B{r}").number_format = MONEY        # IBNR = ultimate − acumulado
ws.cell(33, 5, "=SUM(E23:E32)").number_format = MONEY         # total IBNR (fórmula)

# 3d · comparativo CL vs BF
for j in range(AN):
    r = 37 + j
    ws.cell(r, 1, anios[j])
    ws.cell(r, 2, f"=E{23 + j}").number_format = MONEY        # IBNR CL (referencia a la tabla)
    ws.cell(r, 3, float(ibnr_bf[j])).number_format = MONEY    # IBNR BF (valor)
    ws.cell(r, 4, f"=C{r}-B{r}").number_format = MONEY        # diferencia
ws.cell(47, 2, "=SUM(B37:B46)").number_format = MONEY
ws.cell(47, 3, "=SUM(C37:C46)").number_format = MONEY
ws.cell(47, 4, "=C47-B47").number_format = MONEY

# 3e · CONTROL: reconciliación Python vs Excel (debe dar ≈ 0)
ws.cell(50, 2, round(IBNR_TOTAL, 2)).number_format = MONEY    # Python
ws.cell(51, 2, "=E33").number_format = MONEY                  # Excel (fórmula)
ws.cell(52, 2, "=B51-B50").number_format = MONEY              # diferencia

wb.save("reservas_cedula.xlsx")
print("OK -> reservas_cedula.xlsx")
